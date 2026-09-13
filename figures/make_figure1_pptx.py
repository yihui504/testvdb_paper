# -*- coding: utf-8 -*-
"""Figure 1 (TestVDB pipeline hero figure) generator.

Editable source: figures/pipeline-v6.pptx  (canvas 40 x 13.3 cm, aspect 0.33)
Design language: white bg, thin outlines, no big colour blocks; colour only as
secondary encoding (black/white print safe: independence of the two information
sources is carried by POSITION + dashed line style, not by colour).
Terminology mirrors TestVDB.tex sec.3 verbatim.
"""
from pptx import Presentation
from pptx.util import Cm, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.ns import qn
from pptx.enum.dml import MSO_LINE_DASH_STYLE

# ---------------- palette ----------------
C_STAGE_FILL, C_STAGE_LINE = "F4F8FC", "9DB8D6"
C_NODE_FILL, C_NODE_LINE = "FFFFFF", "4472A4"
C_ART_FILL, C_ART_LINE, C_ART_TEXT = "FBF3DF", "C9A96A", "8A6D3B"
C_GATE_FILL, C_GATE_LINE, C_GATE_TEXT = "9DC3E6", "2E75B6", "1F3864"
C_BLD_FILL, C_BLD_LINE = "EAF4EA", "70AD47"
C_AUD_FILL, C_AUD_LINE = "FCEFD9", "E08A2E"
C_E1_FILL, C_E1_LINE = "F2F2F2", "7F7F7F"
C_E2_FILL, C_E2_LINE = "DDEBF7", "2E75B6"
C_BLACK, C_GRAY, C_BLUE, C_PURPLE = "404040", "595959", "0070C0", "7030A0"
MONO = "Consolas"

prs = Presentation()
prs.slide_width = Cm(40)
prs.slide_height = Cm(13.3)
slide = prs.slides.add_slide(prs.slide_layouts[6])
shapes = slide.shapes


def _set_line(sp, hexcolor, w_pt, dash=None):
    sp.line.color.rgb = RGBColor.from_string(hexcolor)
    sp.line.width = Pt(w_pt)
    if dash:
        sp.line.dash_style = dash


def _no_shadow(sp):
    sp.shadow.inherit = False


def box(x, y, w, h, fill, line, lw=0.75, rounded=True, shape=None):
    st = shape if shape else (MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE)
    sp = shapes.add_shape(st, Cm(x), Cm(y), Cm(w), Cm(h))
    if rounded:
        try:
            sp.adjustments[0] = 0.10
        except Exception:
            pass
    sp.fill.solid()
    sp.fill.fore_color.rgb = RGBColor.from_string(fill)
    _set_line(sp, line, lw)
    _no_shadow(sp)
    tf = sp.text_frame
    tf.word_wrap = True
    tf.margin_left = Cm(0.12); tf.margin_right = Cm(0.12)
    tf.margin_top = Cm(0.04); tf.margin_bottom = Cm(0.04)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    return sp


def lines(sp, items, align=PP_ALIGN.CENTER):
    """items: list of (text, size, bold, italic, color[, font])"""
    tf = sp.text_frame
    for i, it in enumerate(items):
        text, size, bold, italic, color = it[0], it[1], it[2], it[3], it[4]
        font = it[5] if len(it) > 5 else None
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_before = Pt(0); p.space_after = Pt(0)
        p.line_spacing = 1.0
        r = p.add_run(); r.text = text
        r.font.size = Pt(size); r.font.bold = bold; r.font.italic = italic
        r.font.color.rgb = RGBColor.from_string(color)
        if font:
            r.font.name = font
    return sp


def tbox(x, y, w, h, items, align=PP_ALIGN.CENTER, fill=None):
    tb = shapes.add_textbox(Cm(x), Cm(y), Cm(w), Cm(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = Cm(0.03); tf.margin_right = Cm(0.03)
    tf.margin_top = Cm(0.0); tf.margin_bottom = Cm(0.0)
    if fill:
        tb.fill.solid(); tb.fill.fore_color.rgb = RGBColor.from_string(fill)
    else:
        tb.fill.background()
    tb.line.fill.background()
    _no_shadow(tb)
    return lines(tb, items, align)


def conn(x1, y1, x2, y2, color=C_BLACK, w=1.0, dash=None, tail=False):
    c = shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Cm(x1), Cm(y1), Cm(x2), Cm(y2))
    _set_line(c, color, w, dash)
    _no_shadow(c)
    if tail:
        ln = c.line._get_or_add_ln()
        te = ln.makeelement(qn('a:tailEnd'), {'type': 'triangle', 'w': 'med', 'len': 'med'})
        ln.append(te)
    return c


def elbow(pts, color=C_BLACK, w=1.0, dash=None):
    """pts: list of (x,y); arrowhead on last segment"""
    for i in range(len(pts) - 1):
        conn(pts[i][0], pts[i][1], pts[i + 1][0], pts[i + 1][1],
             color, w, dash, tail=(i == len(pts) - 2))


# ---------------- top: external input E1 + automation legend ----------------
tbox(24.0, 0.45, 15.5, 0.5,
     [("reading · generation · adjudication by LLM agents · execution & checks deterministic",
       8, False, True, C_GRAY)], align=PP_ALIGN.RIGHT)
e1 = box(2.0, 0.35, 8.6, 1.35, C_E1_FILL, C_E1_LINE, 0.75)
lines(e1, [("API documentation", 10.5, True, False, "000000"),
           ("natural-language prose", 8, False, False, C_GRAY)])
elbow([(6.3, 1.70), (6.3, 1.98), (4.5, 1.98), (4.5, 2.30)])

# ---------------- stage containers ----------------
ST_Y, ST_H = 2.30, 7.60
stages = [(0.5, 8.0, u"\u2460 Behavioral Specification Extraction"),
          (9.0, 8.0, u"\u2461 Test Script Generation"),
          (17.5, 6.6, u"\u2462 Sandboxed Execution"),
          (24.6, 14.9, u"\u2463 Bug Confirmation \u2014 builder / auditor split")]
for sx, sw, title in stages:
    st = box(sx, ST_Y, sw, ST_H, C_STAGE_FILL, C_STAGE_LINE, 0.75)
    lines(st, [(title, 11.5, True, False, "000000")], align=PP_ALIGN.LEFT)
    st.text_frame.vertical_anchor = MSO_ANCHOR.TOP
    st.text_frame.margin_left = Cm(0.3); st.text_frame.margin_top = Cm(0.12)

# ---------------- stage 1 ----------------
t1 = box(1.1, 3.60, 6.8, 1.30, C_NODE_FILL, C_NODE_LINE)
lines(t1, [("knowledge extractor", 10, True, False, "000000"),
           ("Crawl4AI · coverage + version checks", 7.8, False, True, C_GRAY)])
a1 = box(1.1, 5.20, 6.8, 0.85, C_ART_FILL, C_ART_LINE)
lines(a1, [("knowledge.json", 9.5, True, False, C_ART_TEXT, MONO)])
t2 = box(1.1, 6.35, 6.8, 1.30, C_NODE_FILL, C_NODE_LINE)
lines(t2, [("specification extractor", 10, True, False, "000000"),
           ("categorize · evidence-tier · verify source", 7.8, False, True, C_GRAY)])
a2 = box(1.1, 7.95, 6.8, 1.15, C_ART_FILL, C_ART_LINE)
lines(a2, [("specifications.json", 9.5, True, False, C_ART_TEXT, MONO),
           ("typed · tiered · source-verified · level", 7.5, False, True, C_GRAY)])
conn(4.5, 4.90, 4.5, 5.20, tail=True)
conn(4.5, 6.05, 4.5, 6.35, tail=True)
conn(4.5, 7.65, 4.5, 7.95, tail=True)

# ---------------- stage 2 ----------------
g1 = box(9.6, 3.60, 6.8, 1.30, C_NODE_FILL, C_NODE_LINE)
lines(g1, [("strategy registry", 10, True, False, "000000"),
           ("pre-bound: trigger \u2192 strategy", 7.8, False, True, C_GRAY)])
g2 = box(9.6, 5.20, 6.8, 1.30, C_NODE_FILL, C_NODE_LINE)
lines(g2, [("attack agents", 10, True, False, "000000"),
           ("boundary · state · semantic", 8, False, False, C_GRAY)])
g3 = box(9.6, 6.80, 6.8, 1.30, C_NODE_FILL, C_NODE_LINE)
lines(g3, [("scenario construction", 10, True, False, "000000"),
           ("system-level · no-match fallback", 7.8, False, True, C_GRAY)])
gt = box(9.9, 8.30, 6.2, 0.90, C_GATE_FILL, C_GATE_LINE)
lines(gt, [("5 generation gates", 9.5, True, False, C_GATE_TEXT)])
conn(13.0, 4.90, 13.0, 5.20, tail=True)
conn(13.0, 6.50, 13.0, 6.80, tail=True)
conn(13.0, 8.10, 13.0, 8.30, tail=True)

# ---------------- stage 3 ----------------
cy = box(18.0, 4.00, 5.6, 2.00, C_NODE_FILL, C_NODE_LINE, shape=MSO_SHAPE.CAN)
lines(cy, [("Docker-pinned VDBMS", 9.5, True, False, "000000"),
           ("target @ version", 8, False, True, C_GRAY)])
a3 = box(18.0, 6.60, 5.6, 1.15, C_ART_FILL, C_ART_LINE)
lines(a3, [("raw HTTP logs", 9.5, True, False, C_ART_TEXT, MONO),
           ("+ atomic .done markers", 7.5, False, True, C_GRAY)])
conn(20.8, 6.00, 20.8, 6.60, tail=True)

# ---------------- stage 4: builder / chain / auditor / verdict ----------------
b4 = box(25.2, 3.50, 5.9, 3.00, C_BLD_FILL, C_BLD_LINE)
lines(b4, [("evidence builder", 10, True, False, "000000"),
           ("· doc verification", 7.8, False, False, "375623"),
           ("· execution evidence", 7.8, False, False, "375623"),
           ("· contract grounding", 7.8, False, False, "375623"),
           ("· chain trace", 7.8, False, False, "375623"),
           ("· source grounding (grep clone)", 7.8, False, False, "375623")],
      align=PP_ALIGN.LEFT)
b4.text_frame.vertical_anchor = MSO_ANCHOR.TOP
b4.text_frame.margin_left = Cm(0.25); b4.text_frame.margin_top = Cm(0.15)
ch = box(31.55, 4.35, 3.50, 1.30, C_ART_FILL, C_ART_LINE)
lines(ch, [("evidence_chain.json", 8, True, False, C_ART_TEXT, MONO),
           ("five sections", 7.3, False, True, C_GRAY)])
a4 = box(35.45, 3.50, 3.60, 3.00, C_AUD_FILL, C_AUD_LINE)
lines(a4, [("chain auditor", 10, True, False, "000000"),
           ("4 mechanical checks", 7.8, False, False, "7F4F10"),
           ("4 perspectives:", 7.8, False, False, "7F4F10"),
           ("A contract · B physical", 7.8, False, False, "7F4F10"),
           ("C cognition · D source", 7.8, False, False, "7F4F10")],
      align=PP_ALIGN.LEFT)
a4.text_frame.vertical_anchor = MSO_ANCHOR.TOP
a4.text_frame.margin_left = Cm(0.2); a4.text_frame.margin_top = Cm(0.15)
vd = box(29.6, 7.70, 7.4, 1.25, "FFFFFF", C_NODE_LINE, 1.0)
lines(vd, [("verdict: Confirmed / Refuted", 10, True, False, "000000"),
           ("Neutral \u2192 human review", 8, False, False, C_GRAY)])
conn(31.10, 5.00, 31.55, 5.00, tail=True)
conn(35.05, 5.00, 35.45, 5.00, tail=True)
elbow([(38.00, 6.50), (38.00, 8.32), (37.00, 8.32)])
# bounded rebuild loop: auditor -> builder (dashed)
elbow([(35.70, 6.50), (35.70, 7.00), (28.15, 7.00), (28.15, 6.50)],
      color=C_GRAY, w=1.0, dash=MSO_LINE_DASH_STYLE.DASH)
tbox(30.55, 6.82, 2.6, 0.36, [("rebuild \u2264 3", 7.5, False, True, C_GRAY)], fill="FFFFFF")

# ---------------- cross-stage flows (bus routed below the containers) -------
# S1 -> S2 : specifications.json -> strategy registry, label "constraints"
elbow([(4.5, 9.10), (4.5, 10.15), (9.3, 10.15), (9.3, 4.25), (9.6, 4.25)])
tbox(4.9, 10.20, 2.0, 0.36, [("constraints", 8.5, True, False, C_BLUE)])
# S2 -> S3 : gates -> Docker-pinned VDBMS, label "probes"
elbow([(13.0, 9.20), (13.0, 10.15), (17.75, 10.15), (17.75, 5.00), (18.0, 5.00)])
tbox(13.4, 10.20, 1.2, 0.36, [("probes", 8.5, True, False, C_BLUE)])
# S3 -> S4 : raw HTTP logs -> evidence builder, label "outputs"
elbow([(20.8, 7.75), (20.8, 10.15), (24.9, 10.15), (24.9, 5.60), (25.2, 5.60)])
tbox(21.1, 10.20, 1.5, 0.36, [("outputs", 8.5, True, False, C_BLUE)])

# ---------------- E2: implementation source (independent input) ----------------
e2 = box(28.4, 10.50, 7.6, 1.50, C_E2_FILL, C_E2_LINE, 1.5)
lines(e2, [("implementation source", 10.5, True, False, "000000"),
           ("pinned clone · falsification anchor \u2014", 7.8, False, True, C_PURPLE),
           ("independent of the documentation", 7.8, False, True, C_PURPLE)])
# E2 -> evidence builder (source grounding), purple dashed
elbow([(29.0, 10.50), (29.0, 10.05), (26.2, 10.05), (26.2, 6.50)],
      color=C_PURPLE, w=1.1, dash=MSO_LINE_DASH_STYLE.DASH)
tbox(26.45, 9.62, 3.0, 0.36, [("source grounding", 8, True, False, C_PURPLE)],
     align=PP_ALIGN.LEFT)
# E2 -> chain auditor (D perspective), purple dashed
elbow([(35.2, 10.50), (35.2, 10.05), (38.6, 10.05), (38.6, 6.50)],
      color=C_PURPLE, w=1.1, dash=MSO_LINE_DASH_STYLE.DASH)
tbox(35.45, 9.62, 3.0, 0.36, [("D perspective", 8, True, False, C_PURPLE)],
     align=PP_ALIGN.LEFT)

# ---------------- running example ----------------
tbox(0.5, 12.45, 39.0, 0.5,
     [("Running example (Qdrant #10369):  constraint qdrant_state_recommend_001 \u2192 "
       "state-agent scenario (delete\u2013recreate lookup collection at size 8) \u2192 "
       "200 with silently wrong scores vs. required 400 \u2192 source grounding finds the "
       "dimension check absent on the lookup_from path \u2192 Confirmed (maintainer accepted)",
       7.5, False, True, C_GRAY)])

prs.save(r"C:\Users\11428\Desktop\testvdb_paper\figures\pipeline-v6.pptx")
print("saved pipeline-v6.pptx")
