# -*- coding: utf-8 -*-
"""Figure 1 -- iconic hand-drawn style (style F), mimicking the user's reference:
white bg, dashed/dotted rounded group boxes, bold-outline flat icons,
hand-written font captions, thick black arrows, green check / bug pictogram.
Canvas 40 x 13.3 cm (aspect 0.33). Information content identical to pipeline-v6.
"""
from pptx import Presentation
from pptx.util import Cm, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.dml import MSO_LINE_DASH_STYLE as MSO_LINE
from pptx.oxml.ns import qn

INK = RGBColor(0x11, 0x11, 0x11)
GRAY = RGBColor(0x59, 0x59, 0x59)
BLUE = RGBColor(0x1F, 0x4E, 0x79)
GREEN = RGBColor(0x3F, 0xA3, 0x4D)
RED = RGBColor(0xD9, 0x53, 0x4F)
HAND = "Comic Sans MS"
MONO = "Consolas"

prs = Presentation()
prs.slide_width = Cm(40)
prs.slide_height = Cm(13.3)
slide = prs.slides.add_slide(prs.slide_layouts[6])


def _noline(sp):
    sp.line.fill.background()


def _nofill(sp):
    sp.fill.background()


def rect(x, y, w, h, shape=MSO_SHAPE.RECTANGLE, line_w=2.0, line_c=INK,
         fill=None, dash=None):
    sp = slide.shapes.add_shape(shape, Cm(x), Cm(y), Cm(w), Cm(h))
    sp.shadow.inherit = False
    if fill is None:
        _nofill(sp)
    else:
        sp.fill.solid()
        sp.fill.fore_color.rgb = fill
    if line_w == 0:
        _noline(sp)
    else:
        sp.line.color.rgb = line_c
        sp.line.width = Pt(line_w)
        if dash:
            sp.line.dash_style = dash
    sp.text_frame.word_wrap = True
    return sp


def line(x1, y1, x2, y2, w=1.5, color=INK, dash=None, arrow=False):
    cn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
                                    Cm(x1), Cm(y1), Cm(x2), Cm(y2))
    cn.line.color.rgb = color
    cn.line.width = Pt(w)
    if dash:
        cn.line.dash_style = dash
    if arrow:
        ln = cn.line._get_or_add_ln()
        te = ln.makeelement(qn('a:tailEnd'), {'type': 'triangle', 'w': 'med', 'len': 'med'})
        ln.append(te)
    return cn


def elbow(pts, w=1.5, color=INK, dash=None, arrow=False):
    for i in range(len(pts) - 1):
        is_last = (i == len(pts) - 2)
        line(pts[i][0], pts[i][1], pts[i + 1][0], pts[i + 1][1],
             w=w, color=color, dash=dash, arrow=(arrow and is_last))


def tbox(x, y, w, h, runs, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.TOP,
         space=0.0):
    tb = slide.shapes.add_textbox(Cm(x), Cm(y), Cm(w), Cm(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Cm(0.02)
    tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    for i, para in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = space if space else None
        for (txt, size, bold, font, color) in para:
            r = p.add_run()
            r.text = txt
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.name = font
            r.font.color.rgb = color
            rPr = r._r.get_or_add_rPr()
            ea = rPr.makeelement(qn('a:ea'), {'typeface': font})
            rPr.append(ea)
    return tb


def label(x, y, w, text, size=7.0, color=INK, font=HAND, bold=True, lines=None):
    paras = []
    src = lines if lines else [text]
    for ln_ in src:
        paras.append([(ln_, size, bold, font, color)])
    return tbox(x, y, w, 0.34 * len(paras) + 0.1, paras)


# ---------------------------------------------------------------- icons ----
def icon_book(x, y, w, h):
    rect(x, y, w, h, line_w=2.0)
    line(x + w / 2, y, x + w / 2, y + h, w=1.6)
    for i in range(3):
        yy = y + 0.35 + i * 0.42
        line(x + 0.22, yy, x + w / 2 - 0.2, yy, w=1.0)
        line(x + w / 2 + 0.2, yy, x + w - 0.22, yy, w=1.0)


def icon_page(x, y, w, h, marks=None):
    rect(x, y, w, h, line_w=2.0)
    for i in range(3):
        yy = y + 0.42 + i * 0.42
        line(x + 0.24, yy, x + w - 0.24, yy, w=1.1,
             color=GRAY if not marks else INK)
    if marks:
        tbox(x + 0.12, y + 0.18, w - 0.24, h - 0.3,
             [[("\u2717  \u2713", 8, True, HAND, INK)],
              [("\u2713  \u2713", 8, True, HAND, INK)]])


def icon_code(x, y, w, h):
    rect(x, y, w, h, line_w=2.0)
    rect(x, y, w, h * 0.26, line_w=0, fill=BLUE)
    for i, dx in enumerate((0.18, 0.42, 0.66)):
        rect(x + dx, y + h * 0.08, 0.14, 0.14, shape=MSO_SHAPE.OVAL,
             line_w=0, fill=RGBColor(0xFF, 0xFF, 0xFF))
    tbox(x, y + h * 0.34, w, h * 0.6,
         [[("</>", 13, True, MONO, BLUE)]], anchor=MSO_ANCHOR.MIDDLE)


def icon_cyl(x, y, w, h, txt=None):
    sp = rect(x, y, w, h, shape=MSO_SHAPE.CAN, line_w=2.0, fill=RGBColor(0xFF, 0xFF, 0xFF))
    if txt:
        tbox(x, y + h * 0.35, w, h * 0.5, [[(txt, 8, True, HAND, INK)]],
             anchor=MSO_ANCHOR.MIDDLE)
    return sp


def icon_triple(x, y):
    icon_page(x, y, 1.6, 2.0, marks=True)
    icon_cyl(x + 1.05, y + 1.15, 1.15, 1.05)


def icon_person(x, y, w, h):
    d = w * 0.42
    rect(x + (w - d) / 2, y, d, d, shape=MSO_SHAPE.OVAL, line_w=0, fill=GRAY)
    rect(x, y + d * 0.82, w, h - d * 0.82, shape=MSO_SHAPE.ROUNDED_RECTANGLE,
         line_w=0, fill=GRAY)


def icon_chain(x, y):
    for i in range(5):
        rect(x + i * 0.44, y, 0.68, 0.46, shape=MSO_SHAPE.OVAL, line_w=1.75)


def icon_compare(x, y, d):
    rect(x, y, d, d, shape=MSO_SHAPE.OVAL, line_w=2.25, fill=RGBColor(0xFF, 0xFF, 0xFF))
    line(x + d / 2, y + 0.12, x + d / 2, y + d - 0.12, w=1.6)
    tbox(x + 0.1, y + d * 0.28, d / 2 - 0.15, d * 0.44,
         [[("=", 13, True, HAND, INK)]], anchor=MSO_ANCHOR.MIDDLE)
    tbox(x + d / 2 + 0.05, y + d * 0.28, d / 2 - 0.15, d * 0.44,
         [[("\u2260", 13, True, HAND, INK)]], anchor=MSO_ANCHOR.MIDDLE)


def icon_bug(x, y):
    rect(x + 0.28, y + 0.42, 0.85, 1.1, shape=MSO_SHAPE.OVAL, line_w=0, fill=GRAY)
    rect(x + 0.46, y + 0.06, 0.5, 0.44, shape=MSO_SHAPE.OVAL, line_w=0, fill=GRAY)
    for i, yy in enumerate((0.62, 0.95, 1.28)):
        line(x + 0.3, y + yy, x - 0.02, y + yy - 0.14, w=1.2)
        line(x + 1.11, y + yy, x + 1.43, y + yy - 0.14, w=1.2)
    line(x + 0.55, y + 0.1, x + 0.4, y - 0.14, w=1.2)
    line(x + 0.86, y + 0.1, x + 1.01, y - 0.14, w=1.2)
    rect(x + 0.86, y + 1.06, 0.62, 0.62, shape=MSO_SHAPE.OVAL, line_w=0, fill=RED)
    tbox(x + 0.86, y + 1.1, 0.62, 0.55, [[("!", 10, True, HAND, RGBColor(0xFF, 0xFF, 0xFF))]],
         anchor=MSO_ANCHOR.MIDDLE)


DOT = MSO_LINE.ROUND_DOT
DASH = MSO_LINE.DASH

# ------------------------------------------------- row 1: stages 1-3 ------
rect(0.7, 0.8, 6.4, 4.4, shape=MSO_SHAPE.ROUNDED_RECTANGLE, line_w=1.75, dash=DASH)
icon_book(1.3, 1.3, 2.2, 1.9)
label(0.9, 3.4, 3.0, "", lines=["API documentation (E1)"])
line(3.7, 2.25, 4.5, 2.25, w=1.5, arrow=True)
icon_page(4.6, 1.25, 1.7, 2.0)
label(4.2, 3.4, 2.6, "", lines=["behavioral", "specification (S)"])

rect(8.3, 0.8, 6.4, 4.4, shape=MSO_SHAPE.ROUNDED_RECTANGLE, line_w=1.75, dash=DASH)
icon_page(8.9, 1.25, 1.7, 2.0)
label(8.5, 3.4, 2.6, "", lines=["behavioral", "specification (S)"])
line(10.8, 2.25, 11.6, 2.25, w=1.5, arrow=True)
icon_code(11.7, 1.3, 2.2, 1.9)
label(11.4, 3.4, 2.8, "", lines=["test scripts"])

rect(15.9, 0.8, 9.9, 4.4, shape=MSO_SHAPE.ROUNDED_RECTANGLE, line_w=1.75, dash=DASH)
icon_code(16.5, 1.3, 2.2, 1.9)
label(16.2, 3.4, 2.8, "", lines=["test scripts"])
line(18.9, 2.25, 19.7, 2.25, w=1.5, arrow=True)
icon_cyl(19.8, 1.35, 2.0, 1.85, txt="vdb")
label(19.5, 3.4, 2.6, "", lines=["Docker-pinned", "VDBMS"])
line(22.0, 2.25, 22.8, 2.25, w=1.5, arrow=True)
icon_triple(22.9, 1.25)
label(22.5, 3.4, 3.1, "", lines=["evidence triple", "(resp + logs + state)"])

line(7.1, 2.6, 8.2, 2.6, w=2.5, arrow=True)
line(14.7, 2.6, 15.8, 2.6, w=2.5, arrow=True)

# captions row 1
tbox(0.7, 5.25, 6.4, 0.85,
     [[("\u2460 Specification Synthesis", 10.5, True, HAND, INK)],
      [("categorize \u00b7 evidence-tier \u00b7 verify-source", 7.0, False, HAND, GRAY)]])
tbox(8.3, 5.25, 6.4, 0.85,
     [[("\u2461 Test Script Generation", 10.5, True, HAND, INK)],
      [("pre-bound registry \u00b7 5 generation gates", 7.0, False, HAND, GRAY)]])
tbox(15.9, 5.25, 8.0, 0.85,
     [[("\u2462 Evidence Collection", 10.5, True, HAND, INK)],
      [("Docker-pinned VDBMS \u00b7 raw HTTP logs", 7.0, False, HAND, GRAY)]])

# margin note top-right
tbox(26.6, 1.3, 12.9, 2.6,
     [[("LLM agents: read \u00b7 generate \u00b7 adjudicate", 8.5, True, HAND, INK)],
      [("execution & mechanical checks: deterministic", 8.5, False, HAND, GRAY)],
      [("dual sources: docs (E1, top-left) vs", 8.5, False, HAND, GRAY)],
      [("implementation source (bottom-left)", 8.5, False, HAND, GRAY)]])

# ------------------------------------------- row 2: stage 4 + verdicts ----
# ③ -> ④ elbow routed right of the stage-3 caption
elbow([(24.6, 5.2), (24.6, 6.25), (17.5, 6.25), (17.5, 7.05)], w=2.0, arrow=True)
icon_triple(16.6, 7.2)
label(16.1, 9.35, 3.2, "", lines=["evidence triple"])
line(19.5, 8.2, 20.2, 8.2, w=1.8, arrow=True)
icon_person(20.4, 7.35, 1.7, 1.8)
label(19.9, 9.35, 2.8, "", lines=["evidence builder", "(LLM agent)"])
line(22.4, 8.2, 23.1, 8.2, w=1.8, arrow=True)
icon_chain(23.3, 7.95)
label(22.8, 9.35, 3.2, "", lines=["evidence_chain.json"])
tbox(22.8, 9.68, 3.2, 0.34, [[("(five sections)", 6.5, False, HAND, GRAY)]])
line(26.1, 8.2, 26.8, 8.2, w=1.8, arrow=True)
icon_compare(26.9, 7.1, 2.2)
label(26.4, 9.35, 3.2, "", lines=["chain auditor", "4 mechanical checks", "+ perspectives A\u2013D"])

# verdict fan
line(29.15, 7.7, 31.1, 7.35, w=1.8, arrow=True)
line(29.15, 8.2, 34.0, 8.0, w=1.8, arrow=True)
line(29.15, 8.7, 36.9, 8.6, w=1.8, arrow=True)
tbox(30.9, 6.55, 2.4, 0.9, [[("\u2714", 22, True, HAND, GREEN)]], anchor=MSO_ANCHOR.MIDDLE)
label(30.6, 9.35, 3.0, "", lines=["Pass"])
icon_bug(34.0, 7.0)
label(33.5, 9.35, 3.0, "", lines=["Fail \u2192 bug", "confirmed"])
icon_person(37.0, 7.3, 1.5, 1.6)
label(36.4, 9.35, 3.2, "", lines=["Neutral \u2192", "human review"])

# stage-4 caption
tbox(30.5, 6.15, 9.0, 0.5, [[("\u2463 Evidence-Chain Audit", 10.5, True, HAND, INK)]])

# spec-guidance dotted line (from S icon, via stage gap, down to builder)
elbow([(6.4, 3.0), (7.7, 3.0), (7.7, 6.55), (20.9, 6.55), (20.9, 7.3)],
      w=1.4, dash=DOT, arrow=True)
tbox(10.0, 6.16, 7.0, 0.32, [[("specification (S) guides the builder", 7.0, True, HAND, GRAY)]])

# rebuild loop (dotted, auditor -> builder)
elbow([(28.0, 7.05), (28.0, 6.85), (21.7, 6.85), (21.7, 7.3)], w=1.4, dash=DOT, arrow=True)
tbox(23.2, 6.5, 3.0, 0.3, [[("rebuild \u2264 3", 7.0, True, HAND, GRAY)]])

# --------------------------------------- implementation source (bottom) ---
rect(1.2, 7.0, 5.0, 3.1, shape=MSO_SHAPE.ROUNDED_RECTANGLE, line_w=1.75, dash=DASH)
icon_code(2.0, 7.4, 2.2, 1.9)
label(1.5, 9.4, 4.4, "", lines=["implementation source", "(repo, e.g. sqlite)"])
elbow([(6.2, 8.2), (16.4, 8.2)], w=1.5, dash=DOT, arrow=True)
tbox(8.6, 8.3, 5.4, 0.32, [[("source grounding (C)", 7.0, True, HAND, GRAY)]])
elbow([(3.7, 10.1), (3.7, 10.55), (26.2, 10.55), (26.2, 8.85), (27.0, 8.85)],
      w=1.5, dash=DOT, arrow=True)
tbox(11.5, 10.6, 8.0, 0.32, [[("adversarial cross-check against source (D)", 7.0, True, HAND, GRAY)]])

# running example note
tbox(1.2, 11.35, 37.8, 1.5,
     [[("e.g. #10369 \u2014 docs: ", 8.0, False, HAND, INK),
       ("ON CONFLICT DO UPDATE", 7.5, False, MONO, INK),
       (" returns ", 8.0, False, HAND, INK),
       ("UPDATE", 7.5, False, MONO, INK),
       ("; SQLite 3.51 returns ", 8.0, False, HAND, INK),
       ("INSERT", 7.5, False, MONO, INK),
       ("  \u21d2  Fail, bug confirmed & reported.", 8.0, True, HAND, INK)]])

prs.save(r"C:\Users\11428\Desktop\testvdb_paper\figures\pipeline-style-F-iconic.pptx")
print("saved")
