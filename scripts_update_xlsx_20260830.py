# -*- coding: utf-8 -*-
"""2026-08-30 按 yihui504 最新动作更新 phase1_issue_classification.xlsx
- qdrant #9942: open/PENDING_SELF_LABELED -> closed/TP_FIXED_PR (timvisee: implemented in #10324)
- 新增 10 条表外 issue（chroma 7375, meilisearch 6479-6481, qdrant 10368-10373）
"""
import openpyxl
from copy import copy

wb = openpyxl.load_workbook('data/phase1_issue_classification.xlsx')
ws = wb['issues']
hdr = [c.value for c in ws[1]]
IX = {h: i + 1 for i, h in enumerate(hdr)}

rows = list(ws.iter_rows(min_row=2, values_only=True))
vendor_first = {}
for idx, r in enumerate(rows, start=2):
    if r[0]:
        vendor_first.setdefault(str(r[0]), idx)
print('vendor first row:', vendor_first, '| max_row:', ws.max_row)

QREGEX = r'(?i)qdrant\s+version[:\s]*v?(\d+\.\d+(?:\.\d+)?)'
MREGEX = r'(?i)meilisearch\s+version[:\s]*v?(\d+\.\d+(?:\.\d+)?)'
CREGEX = r'(?i)chroma\s+version[:\s]*v?(\d+\.\d+(?:\.\d+)?)'
D_BASIS = 'D 未裁决：报告者自标，待维护者'

# (vendor, number, repo, title, state, state_reason, created, ver, vsrc, cat, group, label, basis)
NEW = [
    ('chroma', 7375, 'chroma-core/chroma',
     '[Bug]:  Concurrent `create_collection` + `delete_collection` on the same collection can lose the created collection (lost update)',
     'open', None, '2026-07-02', '1.5.9', CREGEX,
     'PENDING_SELF_LABELED', 'D', 'unadjudicated',
     'D 未裁决：报告者自标，待维护者（外部 PR #7643 已挂未合并；社区机制分析存在）'),
    ('meilisearch', 6479, 'meilisearch/meilisearch',
     'Typo tolerance `minWordSizeForTypos` settings have no effect on search results',
     'open', None, '2026-06-30', '1.48.3', MREGEX,
     'PENDING_SELF_LABELED', 'D', 'unadjudicated',
     'D 未裁决：报告者自标（MEMBER ManyTheFish 已回应解释 prefix-search 行为，未裁决关闭）'),
    ('meilisearch', 6480, 'meilisearch/meilisearch',
     'Vector search returns 400 `missing_search_hybrid` despite `hybrid` documented as optional',
     'closed', 'completed', '2026-06-30', '1.48.3', MREGEX,
     'FP_BY_DESIGN', 'C', 'FALSE_POSITIVE',
     'C 误报：维护者判定 by-design（MEMBER ManyTheFish：须显式指定 embedder，行为符合预期，文档无需修改；关闭 completed）'),
    ('meilisearch', 6481, 'meilisearch/meilisearch',
     'Embedder `dimensions=4097` accepted by settings API, exceeding documented maximum',
     'closed', 'completed', '2026-06-30', '1.48.3', MREGEX,
     'FP_BY_DESIGN', 'C', 'FALSE_POSITIVE',
     'C 误报：维护者判定 by-design（dureuill：userProvided 源设计上不设维度上限，4096 文档依据不成立；关闭 completed）'),
    ('qdrant', 10368, 'qdrant/qdrant',
     'Snapshot recovery with `priority: "replica"` silently destroys existing data',
     'open', None, '2026-08-29', '1.19.0', QREGEX,
     'PENDING_SELF_LABELED', 'D', 'unadjudicated',
     'D 未裁决：报告者自标，待维护者（无回应）'),
    ('qdrant', 10369, 'qdrant/qdrant',
     '`recommend` bypasses vector dimension validation for `negative` examples from other collections',
     'open', None, '2026-08-29', '1.19.0', QREGEX,
     'PENDING_SELF_LABELED', 'D', 'unadjudicated',
     'D 未裁决：报告者自标，待维护者（外部 PR #10374 已挂未合并）'),
    ('qdrant', 10370, 'qdrant/qdrant',
     '`PATCH /collections/{c}` with `metadata: {}` does not remove metadata, contrary to versioned docs',
     'closed', 'completed', '2026-08-29', '1.19.0', QREGEX,
     'FP_BY_DESIGN', 'C', 'FALSE_POSITIVE',
     'C 误报：维护者判定 by-design（qdrant-cloud-bot：文档已于 #9907 澄清 merge 语义，行为符合预期；关闭 completed。外部 PR #10379 未随关闭合并）'),
    ('qdrant', 10371, 'qdrant/qdrant',
     '`query/groups` without a query returns nondeterministic group members across identical requests',
     'open', None, '2026-08-29', '1.19.0', QREGEX,
     'PENDING_SELF_LABELED', 'D', 'unadjudicated',
     'D 未裁决：报告者自标，待维护者（外部 PR #10378 已挂未合并）'),
    ('qdrant', 10372, 'qdrant/qdrant',
     '`PUT /collections/{c}/index` accepts `field_schema` as a JSON array and creates wrong index',
     'open', None, '2026-08-29', '1.19.0', QREGEX,
     'PENDING_SELF_LABELED', 'D', 'unadjudicated',
     'D 未裁决：报告者自标，待维护者（外部 PR #10377 已挂未合并）'),
    ('qdrant', 10373, 'qdrant/qdrant',
     '`scroll` on a strict-mode collection ignores `max_query_limit` when `limit` is omitted',
     'open', None, '2026-08-29', '1.19.0', QREGEX,
     'PENDING_SELF_LABELED', 'D', 'unadjudicated',
     'D 未裁决：报告者自标，待维护者（无回应）'),
]

# 1) 更新 qdrant #9942
n_updated = 0
for r in range(2, ws.max_row + 1):
    if str(ws.cell(r, IX['vendor']).value) == 'qdrant' and ws.cell(r, IX['number']).value == 9942:
        ws.cell(r, IX['state']).value = 'closed'
        ws.cell(r, IX['state_reason']).value = 'completed'
        ws.cell(r, IX['gt_category']).value = 'TP_FIXED_PR'
        ws.cell(r, IX['group']).value = 'A'
        ws.cell(r, IX['gt_label']).value = 'CONFIRMED'
        ws.cell(r, IX['classification_basis']).value = 'A 真 bug：已被合并 PR 修复（MEMBER timvisee: implemented in #10324；2026-08-26 关闭 completed）'
        n_updated += 1
print('updated #9942:', n_updated)


def insert_vendor_row(vendor_row, vals):
    ws.insert_rows(vendor_row)
    tpl = ws[vendor_row + 1]  # 移动后的下一行（原行）作为样式模板
    for col, h in enumerate(hdr, start=1):
        c = ws.cell(vendor_row, col, vals.get(h))
        c.font = copy(tpl[col - 1].font)
        c.border = copy(tpl[col - 1].border)
        c.fill = copy(tpl[col - 1].fill)
        c.alignment = copy(tpl[col - 1].alignment)


def to_row(item):
    vendor, number, repo, title, state, sreason, created, ver, vsrc, cat, grp, label, basis = item
    return dict(zip(hdr, [vendor, number, repo, title, state, sreason, False,
                          created + 'T00:00:00Z', ver, vsrc, cat, grp, label, basis, None, None]))

# 2) 插入新行（从后往前插避免行号漂移）
weaviate_first = vendor_first['weaviate']
for item in reversed([x for x in NEW if x[0] == 'qdrant']):
    insert_vendor_row(weaviate_first, to_row(item))

q_first_orig = vendor_first['qdrant']
for item in reversed([x for x in NEW if x[0] == 'meilisearch']):
    insert_vendor_row(q_first_orig, to_row(item))

for item in reversed([x for x in NEW if x[0] == 'chroma']):
    insert_vendor_row(2, to_row(item))

wb.save('data/phase1_issue_classification.xlsx')
print('saved. max_row =', ws.max_row)

# 校验
wb2 = openpyxl.load_workbook('data/phase1_issue_classification.xlsx')
ws2 = wb2['issues']
from collections import Counter
cnt = Counter()
vendor_order = []
prev = None
for r in ws2.iter_rows(min_row=2, values_only=True):
    if not r[0]:
        continue
    v = str(r[0])
    if v != prev:
        vendor_order.append(v)
        prev = v
    cnt[(v, r[10])] += 1
print('vendor order:', vendor_order)
for k in sorted(cnt):
    print('  ', k, cnt[k])
print('total rows:', sum(cnt.values()))
