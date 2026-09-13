"""Look up the two crash/panic-producing confirmed bugs (cycle-2 R1 1.4) and the
current ledger column of the reclassified Qdrant #9149 (cycle-2 R2 3.6)."""
import sys

import openpyxl

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

wb = openpyxl.load_workbook("data/phase1_issue_classification.xlsx",
                            data_only=True)
print("sheets:", wb.sheetnames)
ws = wb[wb.sheetnames[0]]
rows = list(ws.iter_rows(values_only=True))
hdr = [str(h) if h is not None else "" for h in rows[0]]
print("columns:", hdr)
print()

idx = {h: i for i, h in enumerate(hdr)}
crash_col = next((h for h in hdr
                  if "crash" in h.lower() or "panic" in h.lower()), None)
print("crash-related column:", crash_col)
print()

hits = []
for r in rows[1:]:
    if r is None or all(v is None for v in r):
        continue
    blob = " ".join(str(v) for v in r if v is not None).lower()
    if "crash" in blob or "panic" in blob:
        hits.append(r)
print(f"rows mentioning crash/panic: {len(hits)}")
for r in hits[:20]:
    short = [str(v)[:48] for v in r if v is not None][:6]
    print("   ", " | ".join(short))

print("\n=== 9149 ===")
for r in rows[1:]:
    if r is None:
        continue
    blob = " ".join(str(v) for v in r if v is not None)
    if "9149" in blob:
        print("   ", " | ".join(str(v)[:70] for v in r if v is not None))
