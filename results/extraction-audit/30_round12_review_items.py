"""Round-12 review items: (1) core three-run flip cases, (2) full-stage strata
recall under forced / convention / joint scoring, (3) per-system split of the
49 unadjudicated submissions from the phase1 ledger xlsx."""
import glob
import json
import sys
from collections import Counter

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

ROOT = (r"c:/Users/11428/Desktop/testvdb_paper/.paperpilot/phase2-rerun/"
        "arms/rq2_3run")
GT = json.load(open(ROOT + "/gt_81.json", encoding="utf-8"))
ADJ_CONFIRM = {"milvus_005", "milvus_013", "milvus_036", "milvus_038",
               "milvus_043", "qdrant_014", "qdrant_015", "qdrant_016",
               "qdrant_027"}


def load(run, base):
    d = {}
    for f in sorted(glob.glob(f"{ROOT}/{base}/{run}/verdicts_batch*.jsonl")):
        for l in open(f, encoding="utf-8"):
            l = l.strip()
            if not l:
                continue
            o = json.loads(l)
            if o.get("defect_id"):
                d[o["defect_id"]] = o.get("verdict")
    return d


def is_c(v):
    return v in ("CONFIRMED", "HUMAN_REVIEW")


# strata layer, recomputed exactly as 27_final_aggregate.py does
mainv = [json.loads(l) for l in open(
    r"c:/Users/11428/Desktop/testvdb_paper/results/extraction-audit/"
    "rebuild_v1/verify_verdicts_main.jsonl", encoding="utf-8")]
LAYER = {}
for x in mainv:
    if x.get("layer"):
        LAYER[x["case"]] = x["layer"]
AUD = json.load(open(r"c:/Users/11428/Desktop/testvdb_paper/results/"
                     "extraction-audit/rebuild_final/assembly_audit.json",
                     encoding="utf-8"))
for r in AUD:
    c = r["case"]
    if c not in LAYER:
        LAYER[c] = ("evidence_absent" if r["rows_kept"] == 0 and r["arch"] == 0
                    else "evidence_present")
LAYER["milvus_038"] = "weak_evidence"

print("== 1. core three-run flip cases ==")
core_runs = [load(f"run{i}", "rerun_v2") for i in (1, 2, 3)]
ids = sorted(set.intersection(*(set(r) for r in core_runs)))
flips = [i for i in ids if len({r[i] for r in core_runs}) > 1]
print(f"disagreement cases: {len(flips)}/{len(ids)}")
for i in flips:
    pat = [core_runs[k][i] for k in range(3)]
    maj_c = sum(is_c(v) for v in pat) >= 2
    print(f"  {i} GT={GT[i]} {pat} -> majority "
          f"{'CONFIRMED' if maj_c else 'not-confirmed'}")

print("\n== 2. full strata under three readings ==")
full_runs = [load(f"run_full{i}", "rerun_v3") for i in (1, 2, 3)]
for lay in ("evidence_present", "weak_evidence", "evidence_absent"):
    t = sorted(i for i in GT if GT[i] == "T" and LAYER.get(i) == lay)
    if not t:
        continue
    conv = [i for i in t if sum(is_c(r[i]) for r in full_runs) >= 2]
    forc = [i for i in t if sum(r[i] == "CONFIRMED" for r in full_runs) >= 2]
    routed = [i for i in conv
              if any(r[i] == "HUMAN_REVIEW" for r in full_runs)]
    joint = [i for i in conv if i not in routed or i in ADJ_CONFIRM]
    r_joint = [i for i in joint if i in routed]
    print(f"{lay} ({len(t)} packs): convention {len(conv)}/{len(t)}="
          f"{len(conv)/len(t):.3f}  forced {len(forc)}/{len(t)}="
          f"{len(forc)/len(t):.3f}  joint {len(joint)}/{len(t)}="
          f"{len(joint)/len(t):.3f}  (routed-in-conv {len(routed)}, "
          f"of which adjudicated-confirmed {len(r_joint)})")
    if lay == "evidence_absent":
        print("  routed-in-evidence-absent:",
              sorted((i, "ADJ_C" if i in ADJ_CONFIRM else "not-upheld")
                     for i in routed))

print("\n== 3. 49 unadjudicated per-system ==")
try:
    from openpyxl import load_workbook
    wb = load_workbook(r"c:/Users/11428/Desktop/testvdb_paper/data/"
                       "phase1_issue_classification.xlsx", read_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    hdr = [str(h).strip().lower() if h else "" for h in rows[0]]
    print("sheet:", wb.sheetnames[0], "cols:", hdr)
    ci = {h: k for k, h in enumerate(hdr)}
    syscol = next(c for c in ("system", "db", "project") if c in ci)
    stcol = next((c for c in ("status", "adjudication", "state",
                              "adjudication_status") if c in ci), None)
    cnt, total = Counter(), Counter()
    unk = Counter()
    for r in rows[1:]:
        if r[ci[syscol]] is None:
            continue
        sysv = str(r[ci[syscol]]).strip().lower()
        st = str(r[ci[stcol]]).strip().lower() if stcol else ""
        total[sysv] += 1
        if any(k in st for k in ("confirmed", "false", "by-design",
                                 "by_design", "not_repro", "not repro")):
            cnt[sysv] += 1
        else:
            unk[sysv] += 1
    print("total rows per system:", dict(total))
    print("adjudicated-looking per system:", dict(cnt))
    print("non-adjudicated-looking per system:", dict(unk))
except Exception as e:
    print("XLSX FAIL:", repr(e))
