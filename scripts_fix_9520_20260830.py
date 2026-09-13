# -*- coding: utf-8 -*-
"""2026-08-30 修正 qdrant #9520 的错标：TP_FIXED_PR/A -> PENDING_SELF_LABELED/D
证据链：PR #9526 (Fixes #9520) closed unmerged 2026-06-27；
PR #9594 (merged) 为 CodeRabbit bot 在无关性能 PR 上的误交叉引用；
PR #10334 (2026-08-25, AI-assisted) 未合并；issue 仍 open，无成员回应。
"""
import openpyxl

wb = openpyxl.load_workbook('data/phase1_issue_classification.xlsx')
ws = wb['issues']
hdr = [c.value for c in ws[1]]
IX = {h: i + 1 for i, h in enumerate(hdr)}

hit = 0
for r in range(2, ws.max_row + 1):
    if (str(ws.cell(r, IX['vendor']).value) == 'qdrant'
            and ws.cell(r, IX['number']).value == 9520):
        ws.cell(r, IX['gt_category']).value = 'PENDING_SELF_LABELED'
        ws.cell(r, IX['group']).value = 'D'
        ws.cell(r, IX['gt_label']).value = 'unadjudicated'
        ws.cell(r, IX['classification_basis']).value = (
            'D 未裁决：报告者自标，待维护者。'
            '【2026-08-30 错标纠正】原 A/TP_FIXED_PR 系误判：所依据的已合并 PR #9594 实为 '
            'CodeRabbit bot 在无关性能 PR（edge shard 并行加载）上的误交叉引用（"shard" 关键词撞车）。'
            '真实修复 PR #9526（Fixes #9520）2026-06-27 被关闭未合并；#10334（2026-08-25，AI-assisted）'
            '亦未合并；issue 仍 open，无成员回应。'
        )
        hit += 1
print('rows fixed:', hit)

# 同步 legend 计数
from collections import Counter
cnt = Counter()
for r in ws.iter_rows(min_row=2, values_only=True):
    if r[0] and r[10]:
        cnt[str(r[10])] += 1

NEW = {
    'TP_FIXED_PR': cnt.get('TP_FIXED_PR', 0),
    'TP_ACK_OPEN': cnt.get('TP_ACK_OPEN', 0),
    'TP_ACK_CLOSED_NOFIX': cnt.get('TP_ACK_CLOSED_NOFIX', 0),
    'TP_DUP_TRACKED': cnt.get('TP_DUP_TRACKED', 0),
    'FP_BY_DESIGN': cnt.get('FP_BY_DESIGN', 0),
    'FP_NOT_REPRO': cnt.get('FP_NOT_REPRO', 0),
    'BY_DESIGN': cnt.get('BY_DESIGN', 0),
    'PENDING_SELF_LABELED': cnt.get('PENDING_SELF_LABELED', 0),
    'SELF_CLOSED': cnt.get('SELF_CLOSED', 0),
    'STALE_NO_FIX': cnt.get('STALE_NO_FIX', 0),
    'OPEN_NO_LABEL': cnt.get('OPEN_NO_LABEL', 0),
    'SELF_PR_CLOSED': cnt.get('SELF_PR_CLOSED', 0),
    'SELF_PR_OPEN': cnt.get('SELF_PR_OPEN', 0),
}
lg = wb['legend']
for r in range(2, lg.max_row + 1):
    cat = lg.cell(r, 3).value
    if cat in NEW:
        lg.cell(r, 5).value = NEW[cat]
    elif lg.cell(r, 1).value and '合计' in str(lg.cell(r, 1).value):
        a = str(lg.cell(r, 1).value)
        if a.startswith('合计 A∪B∪C'):
            lg.cell(r, 5).value = sum(NEW[k] for k in
                ['TP_FIXED_PR', 'TP_ACK_OPEN', 'TP_ACK_CLOSED_NOFIX', 'TP_DUP_TRACKED',
                 'FP_BY_DESIGN', 'FP_NOT_REPRO', 'BY_DESIGN'])
        elif a.startswith('合计 D'):
            lg.cell(r, 5).value = sum(NEW[k] for k in
                ['PENDING_SELF_LABELED', 'SELF_CLOSED', 'STALE_NO_FIX', 'OPEN_NO_LABEL'])

wb.save('data/phase1_issue_classification.xlsx')
scored = sum(NEW[k] for k in ['TP_FIXED_PR', 'TP_ACK_OPEN', 'TP_ACK_CLOSED_NOFIX', 'TP_DUP_TRACKED',
                              'FP_BY_DESIGN', 'FP_NOT_REPRO', 'BY_DESIGN'])
unadj = sum(NEW[k] for k in ['PENDING_SELF_LABELED', 'SELF_CLOSED', 'STALE_NO_FIX', 'OPEN_NO_LABEL'])
print(f"A={NEW['TP_FIXED_PR']} scored={scored} unadj={unadj} total={scored + unadj + NEW['SELF_PR_CLOSED'] + NEW['SELF_PR_OPEN']}")
