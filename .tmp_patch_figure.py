"""Regenerate the pipeline figure with labels matching Section 3.5:
A contract / B physical / C behavioral elegance / D maintainer cognition;
the implementation source as falsification anchor (not a perspective); and the
three-valued verdict Confirmed / False-Positive / Human-Review.
Writes pipeline-v10.pptx and leaves v9 untouched."""
import shutil
import sys
from pptx import Presentation
from pptx.util import Pt

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

SRC = "figures/pipeline-v9.pptx"
DST = "figures/pipeline-v10.pptx"
shutil.copyfile(SRC, DST)

prs = Presentation(DST)
print("slide WxH cm:", prs.slide_width / 360000, prs.slide_height / 360000)
s = prs.slides[0]
shapes = list(s.shapes)


def runs(shape):
    return [r for p in shape.text_frame.paragraphs for r in p.runs]


# (1) perspective legend: C was mislabelled "cognition", D "source"
sh = shapes[27]
p1 = sh.text_frame.paragraphs[1]
p1.runs[1].text = "C behavioral elegance"
p1.runs[2].text = "\x0bD maintainer cognition"
for r in p1.runs:
    r.font.size = Pt(8.5)
print("[27]", sh.text_frame.text.replace("\x0b", " | "))

# (2) the source is the falsification anchor, not perspective D
sh = shapes[41]
ps = sh.text_frame.paragraphs
ps[0].runs[0].text = "falsification anchor"
ps[1].runs[0].text = "source grounding "
print("[41]", sh.text_frame.text.replace("\n", " | "))

# (3)(4) verdict panel -> the paper's three-valued verdict
sh = shapes[33]
runs(sh)[0].text = "False-Positive"
print("[33]", sh.text_frame.text)

sh = shapes[34]
rr = runs(sh)
rr[0].text = "Human-Review"
rr[1].text = ""
rr[2].text = ""
print("[34]", sh.text_frame.text)

prs.save(DST)
print("saved", DST)
