"""Fix the perspective legend line break: python-pptx's run.text setter escapes
the vertical tab, so insert a real <a:br/> element instead."""
import copy
import shutil
import sys
from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Pt

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

shutil.copyfile("figures/pipeline-v9.pptx", "figures/pipeline-v10.pptx")
prs = Presentation("figures/pipeline-v10.pptx")
s = prs.slides[0]
shapes = list(s.shapes)


def runs(shape):
    return [r for p in shape.text_frame.paragraphs for r in p.runs]


# ---- legend: A contract / B physical / C behavioral elegance / D maintainer cognition
sh = shapes[27]
p1 = sh.text_frame.paragraphs[1]
p1.runs[1].text = "C behavioral elegance"
p1.runs[2].text = "D maintainer cognition"
# insert a line break before the D run (which is p1.runs[2])
br = p1._p.makeelement(qn("a:br"), {})
p1.runs[2]._r.addprevious(br)
for r in p1.runs:
    r.font.size = Pt(8)
print("[27]", sh.text_frame.text.replace("\n", " / "))

# ---- the source is the falsification anchor, not perspective D
sh = shapes[41]
ps = sh.text_frame.paragraphs
ps[0].runs[0].text = "falsification anchor"
ps[0].runs[1].text = ":"
ps[1].runs[0].text = "source grounding "
for r in runs(sh):
    r.font.size = Pt(8.5)

# ---- verdict panel -> Confirmed / False-Positive / Human-Review
runs(shapes[33])[0].text = "False-Positive"
rr = runs(shapes[34])
rr[0].text = "Human-Review"
rr[1].text = ""
rr[2].text = ""

prs.save("figures/pipeline-v10.pptx")
print("saved")
