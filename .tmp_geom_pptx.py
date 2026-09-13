import sys
from pptx import Presentation
from pptx.util import Emu
sys.stdout.reconfigure(encoding="utf-8", errors="replace")
prs = Presentation("figures/pipeline-v9.pptx")
s = prs.slides[0]
EMU_CM = 360000
for shi in (27, 41, 31, 33, 34, 40):
    sh = list(s.shapes)[shi]
    print(f"[{shi}] {sh.name}: L={sh.left/EMU_CM:.2f} T={sh.top/EMU_CM:.2f} "
          f"W={sh.width/EMU_CM:.2f} H={sh.height/EMU_CM:.2f} cm")
    tf = sh.text_frame
    tf.word_wrap = tf.word_wrap
    for pi, p in enumerate(tf.paragraphs):
        for ri, r in enumerate(p.runs):
            sz = r.font.size.pt if r.font.size else None
            print(f"     p{pi}r{ri} size={sz} text={r.text!r}")
