import sys
from pptx import Presentation
sys.stdout.reconfigure(encoding="utf-8", errors="replace")
prs = Presentation("figures/pipeline-v9.pptx")
print("slides:", len(prs.slides))
for si, s in enumerate(prs.slides):
    for shi, sh in enumerate(s.shapes):
        if not sh.has_text_frame:
            continue
        txt = sh.text_frame.text
        if any(k in txt for k in ("perspective", "Perspective", "Refuted",
                                  "Neutral", "verdict", "Verdict",
                                  "cognition", "source grounding")):
            print(f"--- slide{si} shape{shi} name={sh.name!r} ---")
            print(repr(txt))
            for pi, p in enumerate(sh.text_frame.paragraphs):
                for ri, r in enumerate(p.runs):
                    print(f"    p{pi} r{ri}: {r.text!r}")
