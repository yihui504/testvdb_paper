# -*- coding: utf-8 -*-
"""更新 legend sheet 计数（2026-08-30 更新后）"""
import openpyxl

wb = openpyxl.load_workbook('data/phase1_issue_classification.xlsx')
ws = wb['issues']
from collections import Counter
cnt = Counter()
for r in ws.iter_rows(min_row=2, values_only=True):
    if r[0] and r[10]:
        cnt[str(r[10])] += 1

NEW_COUNTS = {
    'TP_FIXED_PR': cnt.get('TP_FIXED_PR', 0),
    'TP_ACK_OPEN': cnt.get('TP_ACK_OPEN', 0),
    'TP_ACK_CLOSED_NOFIX': cnt.get('TP_ACK_CLOSED_NOFIX', 0),
    'TP_DUP_TRACKED': cnt.get('TP_DUP_TRACKED', 0),
    'FP_BY_DESIGN': cnt.get('FP_BY_DESIGN', 0),
    'FP_NOT_REPRO': cnt.get('FP_NOT_REPRO', 0),
    'BY_DESIGN': cnt.get('BY_DESIGN', 0),
    'PENDING_SELF_LABELED': cnt.get('PENDING_SELF_LABELED', 0),
    'SELF_CLOSED': cnt.get('SELF_CLOSED', 0),
    'STALE_NO_FIX': cnt.get('STALE_NO_FIX', 0),
    'OPEN_NO_LABEL': cnt.get('OPEN_NO_LABEL', 0),
    'SELF_PR_CLOSED': cnt.get('SELF_PR_CLOSED', 0),
    'SELF_PR_OPEN': cnt.get('SELF_PR_OPEN', 0),
}

lg = wb['legend']
n = 0
for r in range(2, lg.max_row + 1):
    cat = lg.cell(r, 3).value
    if cat in NEW_COUNTS:
        lg.cell(r, 5).value = NEW_COUNTS[cat]
        n += 1
    elif cat is None and lg.cell(r, 1).value and '合计' in str(lg.cell(r, 1).value):
        pass  # 合计行下面单独处理

# 重算合计行
scored = sum(NEW_COUNTS[k] for k in ['TP_FIXED_PR', 'TP_ACK_OPEN', 'TP_ACK_CLOSED_NOFIX', 'TP_DUP_TRACKED',
                                     'FP_BY_DESIGN', 'FP_NOT_REPRO', 'BY_DESIGN'])
unadj = sum(NEW_COUNTS[k] for k in ['PENDING_SELF_LABELED', 'SELF_CLOSED', 'STALE_NO_FIX', 'OPEN_NO_LABEL'])
excluded = sum(NEW_COUNTS[k] for k in ['SELF_PR_CLOSED', 'SELF_PR_OPEN'])
for r in range(2, lg.max_row + 1):
    a = lg.cell(r, 1).value
    if a is None:
        continue
    a = str(a)
    if a.startswith('合计 A∪B∪C'):
        lg.cell(r, 5).value = scored
    elif a.startswith('合计 D'):
        lg.cell(r, 5).value = unadj
    elif a.startswith('EXCLUDED'):
        lg.cell(r, 5).value = excluded

wb.save('data/phase1_issue_classification.xlsx')
print(f'updated {n} category counts; scored={scored} unadj={unadj} excluded={excluded} total={scored+unadj+excluded}')
