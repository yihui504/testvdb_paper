import sys
from pptx import Presentation
sys.stdout.reconfigure(encoding="utf-8", errors="replace")
prs = Presentation("figures/pipeline-v9.pptx")
s = prs.slides[0]
for shi, sh in enumerate(s.shapes):
    if sh.has_text_frame and sh.text_frame.text.strip():
        print(f"[{shi}] {sh.name}: {sh.text_frame.text!r}")
